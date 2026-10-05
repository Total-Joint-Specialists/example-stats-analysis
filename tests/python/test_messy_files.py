"""Python-side checks on the messy files: they open in pandas and openpyxl, the
color-coded revisions are visible to Python, and the reshape lesson's answer
key is reachable in Python as well as R."""

from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "data"
WORKBOOK = DATA / "messy_abstraction_workbook.xlsx"
KEYS = DATA / "answer-keys"


def test_workbook_opens_and_red_fill_marks_exactly_the_revisions():
    wb = openpyxl.load_workbook(WORKBOOK)
    key = pd.read_csv(KEYS / "abstraction_workbook_tidy.csv")
    red = set()
    for ws in wb.worksheets:
        header = [str(c.value).strip() if c.value is not None else None for c in ws[4]]
        col = header.index("Study ID") + 1
        for row in range(5, ws.max_row + 1):
            cell = ws.cell(row=row, column=col)
            if cell.fill.fgColor.rgb == "FFFF9999":
                red.add(cell.value)
    assert red == set(key.loc[key["revised"] == 1, "case_id"])


def test_workbook_reads_in_pandas():
    sheets = pd.read_excel(WORKBOOK, sheet_name=None, header=None, dtype=str)
    assert list(sheets) == ["Site A", "Site B"]
    assert sheets["Site B"].shape[1] == sheets["Site A"].shape[1] + 1

