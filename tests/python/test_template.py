"""The data-collection template validates every field except the study ID."""

from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]


def test_template_validates_every_field_but_the_id():
    wb = openpyxl.load_workbook(ROOT / "templates" / "data-collection-template.xlsx")
    assert wb.sheetnames == ["README", "data", "dictionary", "missing_codes"]
    ws = wb["data"]
    header = {c.column_letter: c.value for c in ws[1]}
    rules = {}
    for dv in ws.data_validations.dataValidation:
        letter = str(dv.sqref).split(":")[0].rstrip("0123456789")
        rules[header[letter]] = dv
    assert set(rules) == set(header.values()) - {"study_id"}
    assert rules["sex"].type == "list"
    assert rules["sex"].formula1 == '"Female,Male,Not recorded"'
    assert (rules["bmi"].type, rules["bmi"].formula1, rules["bmi"].formula2) == ("decimal", "10", "80")
    assert rules["surgery_date"].type == "date"


def test_template_has_no_patient_identifier_columns():
    wb = openpyxl.load_workbook(ROOT / "templates" / "data-collection-template.xlsx")
    header = [c.value.lower() for c in wb["data"][1]]
    for banned in ["name", "mrn", "dob", "birth"]:
        assert not any(banned in h for h in header), banned
